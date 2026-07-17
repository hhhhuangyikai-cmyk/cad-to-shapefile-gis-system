import os
import re
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(ROOT, "docs")
OUT_DIR = os.path.join(ROOT, "docs", "docx")


def set_run_font(run, size=None, bold=False, color=None, font_name="Microsoft YaHei"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=6, line=1.15):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(paragraph, after=0)
    run = paragraph.add_run(text.replace("`", ""))
    set_run_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def configure_document(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name, size, color, bold in [
        ("Heading 1", 16, "1F4D78", True),
        ("Heading 2", 13, "2E74B5", True),
        ("Heading 3", 11.5, "1F4D78", True),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("组件式 GIS 开发课程设计")
    set_run_font(run, size=9, color="666666")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("基于 ArcObjects 的 CAD 转 Shapefile 与 GIS 浏览分析系统")
    set_run_font(run, size=8.5, color="666666")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, after=2)
    run = p.add_run(title)
    set_run_font(run, size=20, bold=True, color="0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, after=10)
    run = p.add_run("C# WinForms / ArcObjects SDK 10.8 / ArcGIS Engine 10.8")
    set_run_font(run, size=10.5, color="555555")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, after=16)
    run = p.add_run("生成时间：" + datetime.now().strftime("%Y-%m-%d"))
    set_run_font(run, size=9.5, color="777777")


def add_note_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, before=3, after=3)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, color="1F3A5F")


def add_code_block(doc, code_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=3, after=3, line=1.0)
    run = p.add_run("\n".join(code_lines))
    set_run_font(run, size=8.5, font_name="Consolas")


def is_table_separator(line):
    stripped = line.strip()
    return stripped.startswith("|") and set(stripped.replace("|", "").replace("-", "").replace(":", "").replace(" ", "")) == set()


def parse_table(lines, start):
    table_lines = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        table_lines.append(lines[i].strip())
        i += 1

    rows = []
    for idx, line in enumerate(table_lines):
        if idx == 1 and is_table_separator(line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    return rows, i


def add_markdown_table(doc, rows):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)

    for r, row in enumerate(rows):
        for c in range(col_count):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            if r == 0:
                set_cell_shading(cell, "E8EEF5")
                set_cell_text(cell, text, bold=True, color="0B2545")
            else:
                set_cell_text(cell, text)
    doc.add_paragraph()


def add_inline_runs(paragraph, text, size=10.5):
    # 支持简单反引号代码片段加粗显示，避免正文中满是反引号。
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, bold=True, color="0B2545")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def add_markdown_to_doc(doc, markdown_text):
    lines = markdown_text.splitlines()
    i = 0
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            rows, next_i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            i = next_i
            continue

        if stripped.startswith("# "):
            # 顶部标题已由封面生成，跳过原 Markdown 标题。
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(stripped[3:])
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(stripped[4:])
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_format(p, after=3)
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            set_paragraph_format(p, after=3)
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if stripped.startswith("题目："):
            add_note_box(doc, stripped)
            i += 1
            continue

        p = doc.add_paragraph()
        set_paragraph_format(p, after=6)
        add_inline_runs(p, stripped)
        i += 1


def build_docx(markdown_path, output_path, title):
    doc = Document()
    configure_document(doc, title)

    with open(markdown_path, "r", encoding="utf-8") as f:
        markdown_text = f.read()
    add_markdown_to_doc(doc, markdown_text)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)


def main():
    report_md = os.path.join(DOCS_DIR, "课程设计报告正文框架.md")
    defense_md = os.path.join(DOCS_DIR, "答辩说明.md")
    build_docx(
        report_md,
        os.path.join(OUT_DIR, "课程设计报告正文框架.docx"),
        "课程设计报告正文框架",
    )
    build_docx(
        defense_md,
        os.path.join(OUT_DIR, "答辩说明.docx"),
        "课程设计答辩说明",
    )


if __name__ == "__main__":
    main()
